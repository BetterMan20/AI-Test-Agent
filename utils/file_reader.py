from __future__ import annotations

from pathlib import Path


class FileReader:
    """读取需求文档，支持 .txt / .md / .docx。"""

    @staticmethod
    def read(path: str) -> str:
        p = Path(path)
        ext = p.suffix.lower()

        if ext == ".docx":
            return FileReader._read_docx(p)

        return p.read_text(encoding="utf-8")

    @staticmethod
    def _read_docx(path: Path) -> str:
        from docx import Document

        doc = Document(str(path))
        parts: list[str] = []

        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                parts.append(text)

        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                line = " | ".join(cells)
                if line.strip(" |"):
                    parts.append(line)

        return "\n".join(parts)

    @staticmethod
    def find_input(directory: str = "input") -> str | None:
        """自动发现 input 目录下的需求文档。"""
        d = Path(directory)
        if not d.exists():
            return None

        for ext in ("*.docx", "*.txt", "*.md"):
            files = sorted(d.glob(ext))
            if files:
                return str(files[0])

        return None
